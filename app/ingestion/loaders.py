
import os
import re
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader
from pptx import Presentation


def load_file(path):
    """Extract readable text and metadata from supported documents."""
    ext = Path(path).suffix.lower()

    if ext == ".pdf":
        reader = PdfReader(path)
        pages = []
        extracted_pages = 0

        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            page_text = page_text.strip()

            if page_text:
                pages.append(f"[Page {page_number}]\n{page_text}")
                extracted_pages += 1

        text = "\n\n".join(pages).strip()

        if not text:
            raise ValueError(
                "No readable text was extracted from this PDF. "
                "It may be a scanned document. OCR is required "
                "to read image-only PDF pages."
            )

        return text, {
            "pages": len(reader.pages),
            "extracted_pages": extracted_pages,
            "characters": len(text),
        }

    if ext == ".pptx":
        presentation = Presentation(path)
        slides = []

        for slide_number, slide in enumerate(
            presentation.slides, start=1
        ):
            parts = []

            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = shape.text.strip()
                    if value:
                        parts.append(value)

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        values = [
                            cell.text.strip()
                            for cell in row.cells
                        ]
                        if any(values):
                            parts.append(" | ".join(values))

            if parts:
                slides.append(
                    f"[Slide {slide_number}]\n"
                    + "\n".join(parts)
                )

        text = "\n\n".join(slides).strip()

        if not text:
            raise ValueError(
                "No readable text was found in this PowerPoint."
            )

        return text, {
            "slides": len(presentation.slides),
            "characters": len(text),
        }

    if ext in (".txt", ".md"):
        text = Path(path).read_text(
            encoding="utf-8",
            errors="replace",
        ).strip()

        if not text:
            raise ValueError("The selected file is empty.")

        return text, {"characters": len(text)}

    raise ValueError(
        f"Unsupported file type: {ext}. "
        "Supported formats: PDF, PPTX, TXT, and MD."
    )


def load_url(url):
    """Extract readable text from a single web page."""
    url = (url or "").strip()
    parsed = urlparse(url)

    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise ValueError(
            "Enter a valid web URL beginning with http:// or https://."
        )

    response = requests.get(
        url,
        timeout=20,
        headers={"User-Agent": "NotebookLM-Clone/1.0"},
    )
    response.raise_for_status()

    content_type = response.headers.get("Content-Type", "")
    if "html" not in content_type.lower():
        raise ValueError(
            "This URL did not return an HTML page."
        )

    soup = BeautifulSoup(response.text, "html.parser")

    for tag in soup([
        "script", "style", "noscript", "nav",
        "footer", "header", "svg",
    ]):
        tag.decompose()

    title = (
        soup.title.get_text(" ", strip=True)
        if soup.title
        else url
    )

    text = re.sub(
        r"\s+",
        " ",
        soup.get_text(" ", strip=True),
    ).strip()

    if len(text) < 30:
        raise ValueError(
            "The web page contains too little readable text. "
            "It may require JavaScript or block automated access."
        )

    return text, {"url": url, "title": title}
